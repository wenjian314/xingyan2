"""
星言 - Morse Code Translator
支持：对照表管理、文本⇄迷语互译、中文拼音⇄迷语
"""
import os
import json
import re
from kivy.app import App
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.gridlayout import GridLayout
from kivy.uix.scrollview import ScrollView
from kivy.uix.label import Label
from kivy.uix.textinput import TextInput
from kivy.uix.button import Button
from kivy.uix.spinner import Spinner
from kivy.uix.popup import Popup
from kivy.uix.filechooser import FileChooserListView
from kivy.clock import Clock
from kivy.utils import get_color_from_hex
from kivy.lang import Builder
from kivy.core.clipboard import Clipboard

from pypinyin import pinyin, Style

# ---------- 默认对照表（内嵌，防止无文件时不可用）----------
DEFAULT_TABLE = {
    "A": "迷~", "B": "~迷迷迷", "C": "~迷~迷", "D": "~迷迷", "E": "迷",
    "F": "迷迷~迷", "G": "~~迷", "H": "迷迷迷迷", "I": "迷迷", "J": "迷~~~",
    "K": "~迷~", "L": "迷~迷迷", "M": "~~", "N": "~迷", "O": "~~~",
    "P": "迷~~迷", "Q": "~~迷~", "R": "迷~迷", "S": "迷迷迷", "T": "~",
    "U": "迷迷~", "V": "迷迷迷~", "W": "迷~~", "X": "~迷迷~", "Y": "~迷~~",
    "Z": "~~迷迷",
    "0": "~~~~~", "1": "迷~~~~", "2": "迷迷~~~", "3": "迷迷迷~~",
    "4": "迷迷迷迷~", "5": "迷迷迷迷迷", "6": "~迷迷迷迷", "7": "~~迷迷迷",
    "8": "~~~迷迷", "9": "~~~~迷",
    ".": "迷~迷~迷~", ",": "~~迷迷~~", "?": "迷迷~~迷迷", "'": "迷~~~~迷",
    "!": "~迷~迷~~", "/": "~迷迷~迷", "(": "~迷~~迷", ")": "~迷~~迷~",
    "&": "迷~迷迷迷", ":": "~~~迷迷迷", ";": "~迷~迷~迷", "=": "~迷迷迷~",
    "+": "迷~迷~迷", "~": "~迷迷迷迷~", "_": "迷迷~~迷~", '"': "迷~迷迷~迷",
    "$": "迷迷迷~迷迷~", "@": "迷~~迷~迷", "Ä": "迷~迷~", "Ö": "~~~迷",
    "Ü": "迷迷~~", "ß": "迷迷迷~~迷迷", "É": "迷迷~迷迷", "Ñ": "~~迷~~",
    "Ç": "~迷~迷迷",
}

TABLES_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "tables")
os.makedirs(TABLES_DIR, exist_ok=True)


def save_table(name, mapping):
    path = os.path.join(TABLES_DIR, name + ".json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(mapping, f, ensure_ascii=False, indent=2)
    return path


def load_table_file(path):
    path = str(path)
    if path.endswith(".json"):
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    elif path.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        mapping = {}
        for row in ws.iter_rows(values_only=True):
            if not row or len(row) < 2:
                continue
            ch, code = row[0], row[1]
            if ch is None or code is None:
                continue
            ch = str(ch).strip()
            code = str(code).strip()
            if ch in ("字符", "说明", "十级迷语") or not ch:
                continue
            if re.match(r"^[\u4e00-\u9fff]", ch):
                continue
            mapping[ch.upper()] = code
        return mapping
    else:
        raise ValueError("仅支持 .json 或 .xlsx 文件")


def list_tables():
    names = ["默认对照表"]
    for fn in sorted(os.listdir(TABLES_DIR)):
        if fn.endswith(".json"):
            names.append(fn[:-5])
    return names


def table_path(name):
    if name == "默认对照表":
        return None
    return os.path.join(TABLES_DIR, name + ".json")


class MorseCodec:
    def __init__(self, mapping=None):
        self.set_mapping(DEFAULT_TABLE if mapping is None else mapping)

    def set_mapping(self, mapping):
        self._fwd = {k.upper(): v for k, v in mapping.items()}
        self._rev = {}
        for ch, code in self._fwd.items():
            self._rev[code] = ch

    def to_morse(self, text):
        text = text.upper()
        words = re.split(r"\s+", text)
        out_words = []
        for word in words:
            if not word:
                continue
            codes = []
            for ch in word:
                if ch in self._fwd:
                    codes.append(self._fwd[ch])
                else:
                    codes.append("?")
            out_words.append(" ".join(codes))
        return "  ".join(out_words)

    def from_morse(self, morse):
        morse = morse.strip()
        raw_words = re.split(r"\s{2,}", morse)
        result = []
        for raw in raw_words:
            raw = raw.strip()
            if not raw:
                continue
            chars = raw.split(" ")
            word = "".join(self._rev.get(c.strip(), "?") for c in chars if c.strip())
            result.append(word)
        return " ".join(result)

    def text_to_morse_via_pinyin(self, chinese_text):
        py_list = pinyin(chinese_text, style=Style.NORMAL)
        pinyin_str = " ".join([p[0] for p in py_list])
        return self.to_morse(pinyin_str)


THEME = {
    "bg": "#1a1a2e",
    "card": "#16213e",
    "accent": "#e94560",
    "accent2": "#7b2ff7",
    "text": "#ffffff",
    "subtext": "#aaaaaa",
    "input_bg": "#0f3460",
    "success": "#4caf50",
}


KV = f"""
<Label>:
    color: {THEME['text']}
    font_size: 14

BoxLayout:
    orientation: "vertical"
    padding: 12
    spacing: 10
    canvas.before:
        Color:
            rgba: {THEME['bg']}
        Rectangle:
            pos: self.pos
            size: self.size

    # 顶部标题
    BoxLayout:
        orientation: "horizontal"
        size_hint: (1, None)
        height: 60
        spacing: 10
        Label:
            text: "✨ 星言"
            font_size: 28
            bold: True
            color: {THEME['text']}
            size_hint: (None, 1)
            width: 140
        Label:
            text: "Morse Code Translator"
            font_size: 13
            color: {THEME['subtext']}
            valign: "bottom"

    # 对照表管理栏
    BoxLayout:
        orientation: "horizontal"
        size_hint: (1, None)
        height: 50
        spacing: 8
        Label:
            text: "对照表："
            size_hint: (None, 1)
            width: 70
        Spinner:
            id: table_spinner
            text: "默认对照表"
            size_hint: (0.45, 1)
            background_color: {THEME['input_bg']}
            color: {THEME['text']}
        Button:
            text: "📂 导入"
            size_hint: (0.18, 1)
            background_color: {THEME['accent2']}
            color: {THEME['text']}
            on_press: app.open_file_chooser
        Button:
            text: "➕ 新建"
            size_hint: (0.18, 1)
            background_color: {THEME['accent2']}
            color: {THEME['text']}
            on_press: app.create_new_table

    # 输入区
    Label:
        text: "📝 输入"
        size_hint: (1, None)
        height: 30
        halign: "left"
    TextInput:
        id: input_area
        hint_text: "在此输入文本或迷语..."
        multiline: True
        size_hint: (1, 0.28)
        background_color: {THEME['input_bg']}
        foreground_color: {THEME['text']}
        hint_text_color: {THEME['subtext']}
        font_size: 15
        padding: 10

    # 操作按钮
    GridLayout:
        cols: 3
        size_hint: (1, None)
        height: 50
        spacing: 6
        Button:
            text: "🔤 文本→迷语"
            background_color: {THEME['accent']}
            color: {THEME['text']}
            on_press: app.do_text_to_morse
        Button:
            text: "🔠 迷语→文本"
            background_color: {THEME['accent2']}
            color: {THEME['text']}
            on_press: app.do_morse_to_text
        Button:
            text: "🈶 中文→迷语"
            background_color: "#ff6b6b"
            color: {THEME['text']}
            on_press: app.do_chinese_to_morse

    BoxLayout:
        orientation: "horizontal"
        size_hint: (1, None)
        height: 40
        spacing: 6
        Button:
            text: "🔄 交换"
            size_hint: (0.25, 1)
            background_color: {THEME['card']}
            color: {THEME['text']}
            on_press: app.swap_io
        Button:
            text: "🗑️ 清空"
            size_hint: (0.25, 1)
            background_color: {THEME['card']}
            color: {THEME['text']}
            on_press: app.clear_all
        Button:
            text: "📋 复制结果"
            size_hint: (0.5, 1)
            background_color: {THEME['success']}
            color: {THEME['text']}
            on_press: app.copy_output

    # 输出区
    Label:
        text: "📤 输出"
        size_hint: (1, None)
        height: 30
        halign: "left"
    TextInput:
        id: output_area
        hint_text: "翻译结果将显示在这里..."
        multiline: True
        readonly: True
        size_hint: (1, 0.28)
        background_color: {THEME['card']}
        foreground_color: {THEME['text']}
        font_size: 15
        padding: 10
"""


class XingYanApp(App):
    def build(self):
        self.title = "星言"
        self.icon = self._find_icon()
        self.codec = MorseCodec()
        self.current_table_name = "默认对照表"

        root = Builder.load_string(KV)
        self.root = root

        self.spinner = root.ids.table_spinner
        self.spinner.values = list_tables()
        self.spinner.text = "默认对照表"
        self.spinner.bind(text=self.on_table_change)

        self.input_area = root.ids.input_area
        self.output_area = root.ids.output_area

        return root

    def _find_icon(self):
        for name in ("icon.png", "icon.jpg", "icon.jpeg"):
            p = os.path.join(os.path.dirname(os.path.abspath(__file__)), name)
            if os.path.exists(p):
                return p
        return None

    def on_table_change(self, spinner, text):
        self.current_table_name = text
        if text == "默认对照表":
            self.codec.set_mapping(DEFAULT_TABLE)
        else:
            path = table_path(text)
            if path and os.path.exists(path):
                with open(path, "r", encoding="utf-8") as f:
                    self.codec.set_mapping(json.load(f))
        self._snackbar(f"已切换到：{text}")

    def open_file_chooser(self, *args):
        content = BoxLayout(orientation="vertical")
        fc = FileChooserListView(filters=["*.json", "*.xlsx"])
        content.add_widget(fc)

        btns = BoxLayout(size_hint=(1, None), height=50)
        cancel = Button(text="取消")
        confirm = Button(text="导入")
        btns.add_widget(cancel)
        btns.add_widget(confirm)
        content.add_widget(btns)

        popup = Popup(title="选择对照表文件 (.json / .xlsx)", content=content, size_hint=(0.9, 0.9))

        def do_import(*_):
            sel = fc.selection
            if not sel:
                self._snackbar("请先选择一个文件")
                return
            path = sel[0]
            try:
                mapping = load_table_file(path)
                base = os.path.splitext(os.path.basename(path))[0]
                save_table(base, mapping)
                self.spinner.values = list_tables()
                self.spinner.text = base
                self._snackbar(f"已导入：{base}（{len(mapping)} 条）")
            except Exception as e:
                self._snackbar(f"导入失败：{e}")
            popup.dismiss()

        cancel.bind(on_press=popup.dismiss)
        confirm.bind(on_press=do_import)
        popup.open()

    def create_new_table(self, *args):
        content = BoxLayout(orientation="vertical", padding=10, spacing=10)
        name_input = TextInput(hint_text="对照表名称", size_hint=(1, None), height=44)
        content.add_widget(Label(text="输入新对照表名称：", size_hint=(1, None), height=30))
        content.add_widget(name_input)
        btn = Button(text="创建", size_hint=(1, None), height=44)
        content.add_widget(btn)
        popup = Popup(title="新建对照表", content=content, size_hint=(0.8, 0.5))

        def do_create(*_):
            name = name_input.text.strip()
            if not name:
                self._snackbar("名称不能为空")
                return
            save_table(name, {})
            self.spinner.values = list_tables()
            self.spinner.text = name
            self._snackbar(f"已创建：{name}")
            popup.dismiss()

        btn.bind(on_press=do_create)
        popup.open()

    def do_text_to_morse(self, *args):
        text = self.input_area.text
        if not text.strip():
            self._snackbar("请输入内容")
            return
        self.output_area.text = self.codec.to_morse(text)

    def do_morse_to_text(self, *args):
        morse = self.input_area.text
        if not morse.strip():
            self._snackbar("请输入迷语")
            return
        self.output_area.text = self.codec.from_morse(morse)

    def do_chinese_to_morse(self, *args):
        text = self.input_area.text
        if not text.strip():
            self._snackbar("请输入中文")
            return
        chinese_parts = re.findall(r"[\u4e00-\u9fff]+", text)
        if not chinese_parts:
            self._snackbar("未检测到中文字符")
            return
        combined = "  ".join(self.codec.text_to_morse_via_pinyin(p) for p in chinese_parts)
        self.output_area.text = combined

    def swap_io(self, *args):
        tmp = self.input_area.text
        self.input_area.text = self.output_area.text
        self.output_area.text = tmp

    def clear_all(self, *args):
        self.input_area.text = ""
        self.output_area.text = ""

    def copy_output(self, *args):
        Clipboard.copy(self.output_area.text)
        self._snackbar("已复制到剪贴板")

    def _snackbar(self, msg):
        popup = Popup(
            title="提示",
            content=Label(text=msg, color=THEME['text']),
            size_hint=(0.7, 0.3),
        )
        popup.open()
        Clock.schedule_once(lambda dt: popup.dismiss(), 1.8)


if __name__ == "__main__":
    XingYanApp().run()
